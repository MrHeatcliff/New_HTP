"""Rebuild the Excel workbook from adjacent CSV files only."""
import csv
import hashlib
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

HERE=Path(__file__).resolve().parent
NUMERIC={'Seed','Warmup','Hours','Final score','HNS','seed','warmup_actions',
         'aux_stop_actions','agent_actions','final_eval_episodes','episode_sd',
         'random_reference','human_reference','train_hours','eval_hours',
         'invocation_start_unix','invocation_end_unix'}

def main():
    wb=Workbook();about=wb.active;about.title='Read me'
    descriptions=[
      ('Source tab','gid=1403460040; snapshot preserved in Template original'),
      ('Template match','0 exact matches: 42 original plans use Breakout/size25m; our sweep uses 4 other games/size12m/seed0.'),
      ('Completed','212 runs = 53 arms x 4 games x seed0; 100 final evaluation episodes at 100K executed actions.'),
      ('Filled tracker','Original 42 planned rows marked Not matched + 212 completed runs; empty template slots populated first.'),
      ('Final score','Mean raw return across 100 final evaluation episodes. Not best-checkpoint score.'),
      ('HNS','(Final score - random)/(human - random), ratio, human=1. No clipping.'),
      ('Hours / Start / End','Successful worker invocation elapsed, including phases executed in that invocation. Not training-only or GPU-hours. Prior failed attempts and scheduler queue excluded.'),
      ('Train / Eval hours','Measured successful training phase / sum of ten eval checkpoint phases; see Run details. Do not interpret sum of shared-GPU process hours as allocated GPU-hours.'),
      ('Timezone','Asia/Jakarta (UTC+7). Date display is minute precision; Hours uses full timestamps.'),
      ('W&B URL','Blank: this sweep logged locally, no W&B link was recorded.'),
      ('Commands','Completed rows use portable commands derived from actual train commands, replacing interpreter and logdir. Require original commit 5ee4f27 + source.patch; not arbitrary main branch code. Choose a fresh output path.'),
      ('Scope','1 seed only. Episode SD in Run details is not SD across training seeds. No result is assigned to unmatched template rows.'),
    ]
    for row in descriptions:about.append(row)
    about.column_dimensions['A'].width=26;about.column_dimensions['B'].width=120
    for row in about:
        row[0].font=Font(bold=True)
        row[1].alignment=Alignment(wrap_text=True,vertical='top')
        about.row_dimensions[row[0].row].height=48
    sources=[('Filled tracker','filled_run_tracker.csv'),('Completed runs','completed_runs.csv'),
             ('Template audit','template_audit.csv'),('Template original','template_original.csv'),
             ('Run details','run_details.csv')]
    for title,name in sources:
        ws=wb.create_sheet(title)
        with (HERE/name).open(newline='') as f:rows=list(csv.reader(f))
        headers=rows[0]
        for i,row in enumerate(rows):
            converted=[]
            for j,value in enumerate(row):
                if i and value and headers[j] in NUMERIC:
                    try:value=float(value)
                    except ValueError:pass
                converted.append(value)
            ws.append(converted)
            for cell in ws[ws.max_row]:
                # CSV data is literal, never spreadsheet formulas.
                if isinstance(cell.value,str):cell.data_type='s'
        ws.freeze_panes='C2';ws.auto_filter.ref=ws.dimensions
        for cell in ws[1]:
            cell.fill=PatternFill('solid',fgColor='203B5C');cell.font=Font(color='FFFFFF',bold=True)
            cell.alignment=Alignment(wrap_text=True,vertical='center')
        ws.row_dimensions[1].height=42
        for j,head in enumerate(headers,1):
            ws.column_dimensions[get_column_letter(j)].width=70 if head.startswith(('Notes','Command')) else 24
            if head in ('Hours','Final score','HNS','train_hours','eval_hours','episode_sd'):
                for cells in ws.iter_rows(min_row=2,min_col=j,max_col=j):cells[0].number_format='0.0000'
        if 'Status' in headers:
            col=headers.index('Status')+1
            for cells in ws.iter_rows(min_row=2,min_col=col,max_col=col):
                cell=cells[0];cell.fill=PatternFill('solid',fgColor='E2F0D9' if cell.value=='Completed' else 'FFF2CC')
    target=HERE/'ablation_run_tracker_filled.xlsx';wb.save(target)
    loaded=load_workbook(target,read_only=True,data_only=True)
    assert loaded['Completed runs'].max_row==213 and loaded['Filled tracker'].max_row==255
    assert loaded['Template original'].max_row==58
    loaded.close()
    with (HERE/'checksums.csv').open('w',newline='') as f:
        writer=csv.writer(f,lineterminator='\n');writer.writerow(['file','sha256'])
        for path in sorted(HERE.iterdir()):
            if path.is_file() and path.name!='checksums.csv':
                writer.writerow([path.name,hashlib.sha256(path.read_bytes()).hexdigest()])
    print('XLSX PASS',target,flush=True)

if __name__=='__main__':main()
